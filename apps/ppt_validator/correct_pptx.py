"""FROZEN — Legacy PPT auto-correction script.
Do not modify. New fix logic lives in agents/fix_agent.py and services/fix_service.py.
"""
from __future__ import annotations
import sys
from pathlib import Path


def apply_corrections(input_path: str, output_path: str, findings: list[dict]) -> None:
    """Apply AUTO_FIX findings to a PPTX file (legacy stub)."""
    from pptx import Presentation
    prs = Presentation(input_path)
    for finding in findings:
        if finding.get("action_type") != "AUTO_FIX":
            continue
        _apply_finding(prs, finding)
    prs.save(output_path)


def _apply_finding(prs, finding: dict) -> None:
    slide_n = finding.get("slide", 0)
    if not slide_n or slide_n > len(prs.slides):
        return
    slide = prs.slides[slide_n - 1]
    change_type = finding.get("change_type", "")
    if change_type == "font_replace":
        _replace_font(slide, finding.get("before_value", ""), finding.get("after_value", ""))


def _replace_font(slide, old_font: str, new_font: str) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.font.name == old_font:
                    run.font.name = new_font


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python correct_pptx.py <input.pptx> <output.pptx>")
        sys.exit(1)
    apply_corrections(sys.argv[1], sys.argv[2], [])
