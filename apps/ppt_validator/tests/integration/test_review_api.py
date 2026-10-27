"""Integration tests for the review API."""
import io
import os
import pytest

os.environ.setdefault("LLM_PROVIDER", "mock")
os.environ.setdefault("DATABASE_URL", "sqlite:///test_platform.db")

from fastapi.testclient import TestClient
from pptx import Presentation


def _make_pptx_bytes() -> bytes:
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for title in ["Introduction", "Problem", "Solution", "Timeline", "Team", "Next Steps"]:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def client():
    from db.migrate import run as run_migrations
    run_migrations()
    from app_v2 import app
    return TestClient(app, raise_server_exceptions=False)


def test_health(client):
    resp = client.get("/api/v2/health")
    assert resp.status_code == 200
    assert resp.json()["status"] == "ok"


def test_list_agents(client):
    resp = client.get("/api/v2/agents")
    assert resp.status_code == 200
    agent_ids = [a["agent_id"] for a in resp.json()]
    assert "branding" in agent_ids
    assert "template" in agent_ids


def test_review_run(client):
    pptx_bytes = _make_pptx_bytes()
    resp = client.post(
        "/api/v2/review",
        files={"file": ("test.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        data={"profile": "{}"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert "run_id" in body
    assert "score" in body
    assert "status" in body


def test_get_run(client):
    pptx_bytes = _make_pptx_bytes()
    post = client.post(
        "/api/v2/review",
        files={"file": ("test.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        data={"profile": "{}"},
    )
    run_id = post.json()["run_id"]
    resp = client.get(f"/api/v2/review/{run_id}")
    assert resp.status_code == 200
    assert resp.json()["run_id"] == run_id
