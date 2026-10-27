"""Unit tests for PPTService."""
import io
import pytest
from pptx import Presentation
from pptx.util import Inches
from services.ppt_service import PPTService


def _make_pptx(titles: list[str]) -> io.BytesIO:
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for title in titles:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def test_extract_slide_count(tmp_path):
    buf = _make_pptx(["Slide One", "Slide Two", "Slide Three"])
    path = tmp_path / "test.pptx"
    path.write_bytes(buf.read())
    svc = PPTService()
    data = svc.extract(str(path))
    assert data.slide_count == 3
    assert len(data.slides) == 3


def test_extract_titles(tmp_path):
    buf = _make_pptx(["Introduction", "Solution"])
    path = tmp_path / "test.pptx"
    path.write_bytes(buf.read())
    svc = PPTService()
    data = svc.extract(str(path))
    titles = [s.title for s in data.slides]
    assert "Introduction" in titles
    assert "Solution" in titles


def test_file_hash_is_stable(tmp_path):
    buf = _make_pptx(["Only Slide"])
    content = buf.read()
    path = tmp_path / "test.pptx"
    path.write_bytes(content)
    svc = PPTService()
    d1 = svc.extract(str(path))
    d2 = svc.extract(str(path))
    assert d1.file_hash == d2.file_hash
